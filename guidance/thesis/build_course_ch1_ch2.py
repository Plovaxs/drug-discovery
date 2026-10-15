"""Builds a standalone Chapter I + Chapter II draft (.docx -> .pdf) for the Research Methodology in
Computer Science course assignment (Week 5), using the actual coupled-guidance SBDD research project
as its subject matter. This is a course deliverable, NOT the full S2 thesis build (see build_thesis.py
for that) -- it is deliberately simpler (no front-matter, no TOC page-number resolution) and is written
fresh in English to match the course's own slide deck, pulling facts from TASKS.md and the existing
thesis chapters rather than re-deriving them.

Two-pass citation numbering: pass 1 ("dry") walks build_body() purely to collect which reference keys
are cited (nothing is written to a document); pass 2 renders the real .docx with final, alphabetically
assigned citation numbers already known, so no paragraph ever needs to be rebuilt after the fact.

Run:  python3 guidance/thesis/build_course_ch1_ch2.py
Output: guidance/thesis/Week5_Chapter_I_II_Draft.docx and .pdf (via LibreOffice headless)
"""
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from refs import REFS as BASE_REFS  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DOCX = os.path.join(HERE, 'Week5_Chapter_I_II_Draft.docx')
FONT = 'Times New Roman'

# ------------------------------------------------------------------------------------------------- new references
# Newer sources surfaced during this project's literature audits (2026-10-10 .. 2026-10-14), not yet in refs.py.
# 'v'='kn' (from the project's own search log, not individually re-verified for this draft) -- marked with * below.
NEW_REFS = {
    'bokdiff2025': dict(s=('Khodabandeh Yalabadi', 2025), v='kn',
        t='N. Khodabandeh Yalabadi, M. Yazdani-Jahromi, and O. Garibay, “BoKDiff: Best-of-K and best-of-N alignment for diffusion-based structure-based drug design,” arXiv:2501.15631, 2025 (preprint).'),
    'rayka2025': dict(s=('Rayka', 2025), v='kn',
        t='Rayka et al., “Benchmarking uncertainty quantification methods for protein–ligand binding affinity prediction on Leak-Proof PDBBind,” *Sci. Rep.*, 2025.'),
    'rayka2024': dict(s=('Rayka', 2024), v='kn',
        t='Rayka et al., “ENS-Score: Conformal prediction for protein–ligand binding affinity,” *Mol. Inform.*, 2024.'),
    'parks2020': dict(s=('Parks', 2020), v='kn',
        t='C. D. Parks et al., “Conformal prediction applied to the D3R Grand Challenge binding affinity benchmarks,” *Front. Mol. Biosci.*, 2020.'),
    'zhou2026': dict(s=('Zhou', 2026), v='kn',
        t='Y. Zhou, B. Xue, Y. Zhang, and X. Fu, “Quantifying and preserving equivariance under low-bit weight quantization for 3D neural networks,” arXiv:2601.02213, 2026 (preprint).'),
    'rucker2007': dict(s=('Rücker', 2007), v='kn',
        t='C. Rücker, G. Rücker, and M. Meringer, “y-Randomization and its variants in QSPR/QSAR,” *J. Chem. Inf. Model.*, vol. 47, no. 6, pp. 2345–2357, 2007.'),
    'vantilborg2022': dict(s=('van Tilborg', 2022), v='kn',
        t='D. van Tilborg, A. Alenicheva, and F. Grisoni, “Exposing the limitations of molecular machine learning with activity cliffs,” *J. Chem. Inf. Model.*, vol. 62, no. 23, pp. 5938–5951, 2022.'),
    'kwapien2022': dict(s=('Kwapien', 2022), v='kn',
        t='K. Kwapien, E. Nittinger, J. He, C. Margreitter, A. Voronov, and C. Tyrchan, “Implications of additivity and nonadditivity for machine learning and deep learning models in drug design,” *ACS Omega / J. Chem. Inf. Model.*, 2022.'),
    'thomas2021': dict(s=('Thomas', 2021), v='kn',
        t='M. Thomas, R. T. Smith, N. M. O’Boyle, C. de Graaf, and A. Bender, “Comparison of structure- and ligand-based scoring functions for deep generative models: a GPCR case study,” *J. Cheminform.*, vol. 13, art. 39, 2021.'),
}
REFS = dict(BASE_REFS)
REFS.update(NEW_REFS)


def ordered_keys(keys):
    uniq = list(dict.fromkeys(keys))
    return sorted(uniq, key=lambda k: (REFS[k]['s'][0].lower().replace('ö', 'o').replace('ü', 'u').replace('í', 'i'), REFS[k]['s'][1], k))


# ------------------------------------------------------------------------------------------------- docx helpers
def set_font(run, size=12, bold=None, italic=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


TOKEN = re.compile(r'(\*\*.+?\*\*|\*[^*\s][^*]*?\*)')


def add_runs(par, text, size=12, bold=False, italic=False):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            r = par.add_run(part[2:-2]); set_font(r, size, True, italic)
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            r = par.add_run(part[1:-1]); set_font(r, size, bold, not italic)
        else:
            r = par.add_run(part); set_font(r, size, bold, italic)


def fmt(par, align=WD_ALIGN_PARAGRAPH.JUSTIFY, spacing=1.5, first=0.3, before=0, after=8):
    pf = par.paragraph_format
    par.alignment = align
    pf.line_spacing = spacing
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if first is not None:
        pf.first_line_indent = Inches(first)
    pf.widow_control = True


# ------------------------------------------------------------------------------------------------- two-pass state
DRY = True
CITED = []
NUMBERING = {}
doc = None


def resolve(text):
    """Replace @{key,key2} citation markers with final bracket numbers (real pass) or just record
    which keys were used (dry pass, where nothing is rendered so the exact substitution doesn't matter)."""
    def repl(m):
        keys = [k.strip() for k in m.group(1).split(',')]
        for k in keys:
            if k not in REFS:
                raise KeyError(f'unknown ref key: {k}')
        if DRY:
            CITED.extend(keys)
            return ''
        nums = sorted(NUMBERING[k] for k in keys)
        return '[' + ','.join(str(n) for n in nums) + ']'
    return re.sub(r'@\{([^}]+)\}', repl, text)


def centered(text, size=12, bold=False, italic=False, before=0, after=6):
    text = resolve(text)
    if DRY:
        return
    p = doc.add_paragraph(); add_runs(p, text, size, bold, italic); fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 1.2, first=None, before=before, after=after)


def h1(text):
    if DRY:
        return
    if doc.paragraphs:
        doc.add_page_break()
    p = doc.add_paragraph(); add_runs(p, text, 16, True); fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 1.2, first=None, before=0, after=18)
    p.paragraph_format.keep_with_next = True


def h2(text):
    if DRY:
        return
    p = doc.add_paragraph(); add_runs(p, text, 13, True); fmt(p, WD_ALIGN_PARAGRAPH.LEFT, 1.2, first=None, before=18, after=8)
    p.paragraph_format.keep_with_next = True


def h3(text):
    if DRY:
        return
    p = doc.add_paragraph(); add_runs(p, text, 12, True, True); fmt(p, WD_ALIGN_PARAGRAPH.LEFT, 1.2, first=None, before=12, after=4)
    p.paragraph_format.keep_with_next = True


def para(text, **kw):
    text = resolve(text)
    if DRY:
        return
    p = doc.add_paragraph(); add_runs(p, text); fmt(p, **kw)


def labeled(label, text):
    text = resolve(text)
    if DRY:
        return
    p = doc.add_paragraph()
    r = p.add_run(label + '  '); set_font(r, 12, True, True)
    add_runs(p, text)
    fmt(p, first=0.5, after=8)


def quote(text):
    text = resolve(text)
    if DRY:
        return
    p = doc.add_paragraph()
    r = p.add_run(text); set_font(r, 12, False, True)
    fmt(p, first=0.5, after=10)


def table(headers, rows, caption=None, widths=None):
    rows = [[resolve(str(c)) for c in row] for row in rows]
    if DRY:
        return
    if caption:
        cp = doc.add_paragraph(); add_runs(cp, caption, 11, True); fmt(cp, WD_ALIGN_PARAGRAPH.LEFT, 1.0, first=None, before=10, after=4)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, htext in enumerate(headers):
        c = t.rows[0].cells[j]; c.text = ''
        p = c.paragraphs[0]; add_runs(p, htext, 10.5, True); fmt(p, WD_ALIGN_PARAGRAPH.LEFT, 1.0, first=None, after=0)
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            c = t.rows[i + 1].cells[j]; c.text = ''
            p = c.paragraphs[0]; add_runs(p, cell_text, 10.5); fmt(p, WD_ALIGN_PARAGRAPH.LEFT, 1.0, first=None, after=0)
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Inches(w)
    doc.add_paragraph()


# ================================================================================================== CONTENT
def build_body():
    # ---------------------------------------------------------------------------------------------- title page
    for _ in range(5):
        if not DRY:
            doc.add_paragraph()
    centered('RESEARCH METHODOLOGY IN COMPUTER SCIENCE', 12, True, after=2)
    centered('Week 5 Assignment — Chapter I & II Draft', 12, after=36)
    centered('Coupled Affinity–Synthesizability Guidance for', 15, True, after=0)
    centered('Target-Conditional Molecular Diffusion in', 15, True, after=0)
    centered('Structure-Based Drug Design:', 15, True, after=0)
    centered('A Falsification-Driven Evaluation', 15, True, after=30)
    centered('(Working draft, built directly from the author’s ongoing M.Sc. thesis research)', 11, italic=True, after=36)
    for _ in range(6):
        if not DRY:
            doc.add_paragraph()
    centered('[Author Name]', 12, after=0)
    centered('[Student Identification Number]', 12, after=24)
    centered('October 2026', 12, after=0)

    # ================================================================================================ CHAPTER I
    h1('CHAPTER I — INTRODUCTION')

    h2('1.1  Background')
    para('Structure-based drug design (SBDD) seeks small molecules that bind a protein of known three-dimensional '
         'structure. Classically, candidates are proposed by medicinal chemists or drawn from large compound '
         'libraries, placed in the binding pocket by docking programs, and ranked by scoring functions such as the '
         'empirical function of AutoDock Vina @{trott2010,eberhardt2021}. In recent years, deep generative models '
         'have been added to this workflow. Pocket-conditioned three-dimensional diffusion models, of which '
         'TargetDiff is a representative example @{guan2023}, learn the joint distribution of atom types and '
         'coordinates of ligands inside a pocket and can produce new molecules directly in the binding site, '
         'following the general principle of denoising diffusion probabilistic models @{ho2020}.')
    para('A generative model that reproduces its training distribution is not, by itself, a drug-design tool: most '
         'sampled molecules bind poorly, cannot be synthesized, or are physically implausible. A natural remedy, '
         'borrowed from image generation, is **guidance**. In classifier guidance, the gradient of an auxiliary '
         'predictor is used to steer each denoising step toward inputs that the predictor scores highly '
         '@{dhariwal2021}. For drug design, the auxiliary predictors are learned models of binding affinity and of '
         'synthesizability, and the appeal is obvious: the generator remains fixed, and only a small, cheaply '
         'trained predictor must be supplied. This idea motivates the research programme that this chapter '
         'introduces — the coupled use of an affinity signal and a synthesizability signal to steer a '
         'target-conditional diffusion model, and, when that fails to help, a systematic, falsification-driven '
         'diagnosis of *why* it fails.')

    h2('1.2  Problem Formulation')

    h3('1.2.1  From Topic to Problem')
    table(
        ['Topic', 'Problem', 'Research gap'],
        [[
            'Guided generative structure-based drug design with pocket-conditioned diffusion models.',
            'It is not established whether learned guidance signals — gradient-based or selection-based — improve '
            'generated molecules on outcomes computed independently of the guidance model itself, as opposed to '
            'improving only the guidance model’s own score.',
            'No published evaluation of learned guidance for a 3D diffusion SBDD generator combines a leakage-safe, '
            'target-disjoint evaluation split with statistically corrected significance testing, an explicit '
            'positive control that demonstrates the evaluation harness can detect a genuine improvement, and '
            'mechanism-level diagnostics that explain a null result rather than merely report it.',
        ]],
        widths=[1.6, 2.4, 2.4],
    )

    h3('1.2.2  Anatomy of the Problem Statement')
    labeled('Context.', 'Pocket-conditioned 3-D diffusion generation for SBDD, specifically TargetDiff trained on '
            'CrossDocked2020 @{guan2023,francoeur2020}, and the practitioners who would rely on a guided variant of '
            'such a model to prioritise synthesis candidates under a limited compute budget.')
    labeled('Current knowledge.', 'Gradient-based affinity guidance of diffusion SBDD generators has been reported '
            'to work: BADGER applies the gradient of a separately trained affinity predictor to a frozen diffusion '
            'generator and reports up to a 60% improvement in binding affinity @{jian2026badger}; DiffGui '
            'integrates affinity and drug-likeness guidance trained jointly with the generator and reports improved '
            'binding affinity together with limited wet-laboratory validation @{hu2025diffgui}. A related, older '
            'literature on 2-D generative chemistry shows that docking-score guidance of a reinforcement-learning '
            'generator (REINVENT) can work @{thomas2021}, which is part of why a positive effect was a reasonable '
            'prior expectation for this project. In parallel, a separate body of work shows that (a) docking scores '
            'used to evaluate such claims are known to be biased toward larger molecules @{xu2022,gao2025}; (b) '
            'learned affinity predictors are frequently inflated by data leakage between training and test '
            'complexes @{volkov2022,graber2025,li2026lp}; and (c) best-of-N selection among unguided diffusion '
            'samples is itself an established, effective technique for structure-based generation, independent of '
            'any gradient mechanism @{bokdiff2025}.')
    labeled('Gap.', 'No source examined in Chapter II states that it reports all four of the following together: a '
            'leakage-safe, target-disjoint split; a multiple-comparison-corrected significance test; a positive or '
            'negative control that shows the evaluation harness is sensitive to a real effect; and an explicit '
            'check that a measured gain in docking score is not simply a gain in molecular size. Most '
            'gradient-guidance papers in SBDD do not report a positive control at all.')
    labeled('Consequence.', 'Without these checks, a reported guidance improvement cannot be distinguished from '
            'three alternative explanations that imply different and partly incompatible next steps for the field: '
            'the predictor may carry no usable signal about the interface (better predictors are needed); the '
            'predictor may be informative but the gradient mechanism may misalign with it (better guidance '
            'mechanisms are needed); or the apparent improvement may be an artefact that a differently shaped '
            'objective could exploit without a genuinely better binder (the guidance paradigm itself needs '
            'reconsideration for this generator family). Resolving which one holds is therefore not a cosmetic '
            'detail of method, it is the difference between three incompatible research agendas.')
    labeled('Direction.', 'This research implements gradient-based guidance on a frozen TargetDiff checkpoint with '
            'a leakage-safe affinity predictor, tests it under a pre-registered, multiply-corrected statistical '
            'protocol, adds an explicit positive control built from a genuine (non-surrogate) affinity signal, and '
            '— whichever result obtains — follows up with mechanism-level diagnostics (gradient-field analysis, '
            'representation ablations, class-prior decomposition) designed to adjudicate among the three candidate '
            'explanations above.')

    h3('1.2.3  Problem Statement')
    quote('“Given a frozen, pretrained, pocket-conditioned diffusion model and learned predictors of '
          'binding affinity and synthesizability that are accurate on a leakage-safe, target-disjoint '
          'held-out split, do gradient-based guidance and post-hoc, selection-based guidance built on these '
          'predictors produce statistically detectable, non-degenerate improvements in independently '
          'computed docking score, physical validity, ligand efficiency, and synthesizability — and, when '
          'they do not, which of three candidate mechanisms (an uninformative guidance signal, a correctly '
          'informative signal misapplied through the guidance mechanism, or an objective whose apparent '
          'gain can be obtained without a genuine improvement) explains the failure?”')

    h3('1.2.4  Feasibility and Significance Check')
    labeled('Significant.', 'Guidance is the central mechanism by which generative SBDD methods are expected to '
            'move from valid-molecule generation to practically useful candidate prioritisation; whether it works, '
            'and why it does or does not, has direct consequences for how the field allocates further modelling '
            'effort.')
    labeled('Researchable.', 'A frozen, publicly documented checkpoint (TargetDiff), a public cross-docked '
            'benchmark (CrossDocked2020), and learned predictors trained on it are all available. The full '
            'experimental programme was executed on a single consumer GPU (4 GB VRAM) under an explicit thermal '
            'budget, using a purpose-built, thermally gated job queue (Chapter III) that makes the work tractable '
            'on the hardware actually available for this degree.')
    labeled('Novel enough.', 'Section 1.6 and Section 2.10 detail five claims that survived an eight-round '
            'literature search as defensible contributions: a factorial ablation separating interface *identity* '
            'from interface *quantity*; a class-prior decomposition of apparent predictive performance, replicated '
            'for the deployed architecture itself; a task-level statistical power analysis; the explicit connection '
            'from an information audit of the predictor to the observed guidance failure; and ternary (1.58-bit) '
            'quantization-aware training of a 3-D equivariant affinity network, validated on both accuracy and '
            'geometric equivariance.')
    labeled('Scoped.', 'Five mechanistically distinct guidance routes — two gradient-mechanism families, one '
            'selection-based route, one synthesizability-guidance route, and one predictor-redesign route — share a '
            'single evaluation harness, a single frozen generator, and one family of predictors, so extending the '
            'investigation from one route to the next required no new infrastructure, only a new predictor or a '
            'new sampling-time hook. This keeps a five-route falsification programme tractable within a master’s '
            'timeline and a 4 GB GPU.')

    h2('1.3  Research Questions and Hypotheses')
    para('The problem statement above is decomposed into six research questions. Each is tagged by its question '
         'form, following the taxonomy of descriptive, comparative, relational, and causal questions. Four of the '
         'six admit a sharp null/alternative hypothesis pair and a pre-registered statistical test; the remaining '
         'two are diagnostic or decompositional and are answered by estimation (a correlation, or a variance '
         'decomposition with a confidence interval) rather than by a single significance test — a distinction the '
         'course materials themselves draw between hypothesis-testing and other legitimate forms of quantitative '
         'inference. Status lines summarise what has already been measured; full statistical detail is deferred to '
         'the methodology and results chapters of the underlying thesis.')

    labeled('RQ1 [Causal — primary].', 'Under a leakage-safe, target-disjoint evaluation protocol, does '
            'gradient-based guidance from a learned affinity predictor, applied to a frozen pocket-conditioned '
            'diffusion generator, produce a statistically significant, non-degenerate improvement in independently '
            'computed docking score relative to an unguided baseline?')
    labeled('   H0 / H1.', 'H0: the median independently computed Vina Dock score of guided samples is not '
            'different from the unguided baseline. H1: guided samples show a significantly more negative (better) '
            'docking score, without a significant degradation in ligand efficiency or physical validity. '
            '*Status:* refuted for the gradient route at every tested guidance strength and under two predictor '
            'architectures and three mechanism variants (no condition produced a significant effect; one mechanism '
            'variant was significantly *worse*).')

    labeled('RQ2 [Causal — harness validation / positive control].', 'Does post-hoc, best-of-N selection using a '
            'genuine, non-circular binding-affinity signal (AutoDock Vina itself, not the trained surrogate) '
            'produce a detectable improvement on the same generator, the same pockets, and the same independent '
            'evaluation pass used for RQ1 — thereby showing that the evaluation harness is capable of detecting an '
            'improvement when a genuinely informative signal is available?')
    labeled('   H0 / H1.', 'H0: no difference between Vina-reranked and unguided baseline samples. H1: '
            'Vina-reranked samples show a significantly better independently computed docking score. *Status:* '
            '**confirmed** on 60 pocket–seed trials (20 pockets × 3 seeds): Vina Dock −8.07 ± 2.66 kcal/mol '
            '(reranked) vs. −6.81 ± 10.08 kcal/mol (baseline), Wilcoxon p = 1.9 × 10⁻⁶, target-clustered bootstrap '
            'CI for the per-pocket delta [−2.15, −0.74] kcal/mol, excluding zero. This directly validates the '
            'evaluation harness and localises the RQ1 failure to the *signal*, not to an insensitive harness — '
            'though a real confound must be disclosed alongside it: reranked molecules are on average 1.8 heavy '
            'atoms larger (p = 0.0002) and less synthetically accessible (SA and RA score both drop significantly, '
            'p < 10⁻⁴), while ligand efficiency stays flat (p = 0.756) and physical validity does not fall.')

    labeled('RQ3 [Relational / diagnostic].', 'What does the gradient of the learned affinity predictor actually '
            'encode with respect to atomic position — is its magnitude related to proximity to the true binding '
            'interface, as a correctly functioning guidance signal would require?')
    labeled('   Status.', 'A factorial ablation that removes exactly 50% of protein atoms under three conditions '
            '(nearest half kept — interface intact; farthest half kept — interface destroyed; random half kept) '
            'shows the predictor is roughly 4.5× more sensitive to *how many* protein atoms are present than to '
            '*whether* the interface is intact (destroying real contacts is, if anything, cheaper than removing '
            'atoms that never touch the ligand). A complementary measurement finds the gradient’s magnitude '
            '*increases* with distance from the pocket in the sampled states — the wrong direction for a signal '
            'that is meant to pull atoms toward contact-forming positions.')

    labeled('RQ4 [Comparative].', 'Does a typed, chemically explicit interaction representation (hydrogen-bond, '
            'hydrophobic, π-stacking and salt-bridge contact counts, extracted with PLIP) provide additional '
            'predictive information about binding affinity beyond ligand-only or pocket-only descriptors, once the '
            'confound of molecular/pocket size is controlled for?')
    labeled('   H0 / H1.', 'H0: no significant difference in target-clustered R² between a model trained on the '
            'typed-interaction block and models trained on ligand-only or pocket-only blocks, after '
            'Benjamini–Hochberg correction. H1: the typed-interaction block yields significantly higher R². '
            '*Status:* **H0 retained** — the four typed-interaction counts, summed crudely per complex, reach only '
            'R² = 0.07–0.08 and lose significantly to both the ligand-only block (R² = 0.353, q = 0.003) and the '
            'pocket-only block (R² = 0.289, q = 0.003); adding a single size covariate (heavy-atom count) alone '
            'raises R² to 0.324 (ΔR² = +0.242, q < 0.001). Almost all of the usable signal in this block is size, '
            'not interaction type.')

    labeled('RQ5 [Descriptive / relational — class-prior decomposition].', 'What fraction of a learned affinity '
            'model’s apparent predictive performance is attributable to a class-level prior (the typical affinity '
            'of the protein family or chemotype class targeted) rather than genuine within-target '
            'structure–activity relationships, and does this fraction hold for the deployed 3-D equivariant '
            'architecture itself, not only for a descriptor-based baseline?')
    labeled('   Status.', 'Measured by within-target label permutation (the per-target label multiset is '
            'preserved, but which ligand carries which label is randomised): a 14-descriptor ridge regression '
            'reaches R² = 0.247 ± 0.002 under this control, i.e. 70% of its real-label performance (0.353) is '
            'reproducible without any within-target signal. The same control, run directly on the deployed EGNN '
            'architecture, reaches R² = 0.307 — **90%** of the EGNN’s own real-label performance (0.342). The '
            'architecture reaches its class prior through a more direct route (reading 3-D target geometry '
            'directly) rather than through chemotype correlation, but the share of apparent performance explained '
            'by class membership, not local structure–activity relationship, is larger for the deployed model than '
            'for the descriptor baseline, not smaller.')

    labeled('RQ6 [Relational].', 'Given the measured ceiling imposed by experimental label noise and the measured '
            'between-seed variance of the deployed architecture, is the window within which two models’ true '
            'performance can be distinguished statistically wider or narrower than a single confidence interval on '
            'this task?')
    labeled('   Status.', 'The window between the floor set by the weakest defensible baseline (ligand-only ridge, '
            'R² = 0.353) and the ceiling set by label noise (R² = 0.578, estimated from repeat-measurement '
            'correlation) spans only 0.225 R², which is *narrower* than the average width of a single '
            'target-clustered bootstrap confidence interval on this task (0.387; ratio 0.58). A single benchmark '
            'comparison reporting a gain of 0.02–0.05 R² on similarly sized data is, by this measure, not able to '
            'resolve where within that window a model actually sits.')

    h2('1.4  Operational Definitions')
    table(
        ['Construct', 'Operational measure', 'Unit of analysis', 'Measurement rule'],
        [
            ['Guidance effectiveness (RQ1, RQ2)',
             'Independently computed AutoDock Vina Dock score (mode = dock, exhaustiveness = 16), plus ligand '
             'efficiency, PoseBusters validity, QED, SA score and RA score as co-reported quality checks.',
             'One generated molecule–pocket trial.',
             'Guidance (gradient or selection) is applied only during sampling or candidate selection; the final '
             'score is always computed by a separate Vina pass with the same settings used for the unguided '
             'baseline, never by the selection-time score itself (avoids circularity).'],
            ['Binding-affinity predictive quality (RQ4, RQ5, RQ6)',
             'Coefficient of determination (R²) between predicted and experimental pK on a held-out split.',
             'Target cluster (for bootstrap confidence intervals); individual protein–ligand complex (for the '
             'underlying point predictions).',
             'Leakage-safe, target-disjoint split (an LP-PDBBind-style adaptation applied to CrossDocked2020); '
             'target-clustered bootstrap, B = 2000; Benjamini–Hochberg correction across all paired comparisons '
             'reported together.'],
            ['Chemical novelty of a test ligand',
             'Maximum ECFP4 Tanimoto similarity to any training-set ligand, binned into four tiers (novel < 0.35, '
             'distant 0.35–0.50, related 0.50–0.70, same series ≥ 0.70).',
             'Individual test-set ligand.',
             'Computed once against the full training pool after the leakage-safe split is fixed; tiers are '
             'defined before any performance analysis is run on them.'],
            ['Class-level prior (RQ5)',
             'R² of a model whose features are unchanged but whose training labels are permuted *within* each '
             'target (so the per-target label multiset, i.e. the class-typical affinity range, is preserved while '
             'ligand-to-label assignment is destroyed).',
             'Protein target (permutation is applied within each target’s rows).',
             '20 permutations for the descriptor pipeline; one full retraining run under a fixed seed for the EGNN '
             'architecture, using the same hyperparameters and split as the deployed checkpoint.'],
        ],
        widths=[1.6, 2.3, 1.1, 2.3],
    )

    h2('1.5  Research Design Overview')
    para('Table 1.2 aligns each research question with the design type used to answer it, following the '
         'question-to-design-to-claim mapping introduced in the course materials. All designs reuse one frozen '
         'generator, one pocket set, and one statistical protocol (target-clustered bootstrap with '
         'Benjamini–Hochberg correction), so that a claim of “no effect” in one route is comparable in strength to '
         'a claim of “effect” in another.')
    table(
        ['RQ', 'Design type', 'Comparison / manipulation', 'Primary outcome'],
        [
            ['RQ1', 'Within-generator repeated-measures experiment (benchmarking, replicated across targets)',
             'Gradient-guided vs. unguided sampling, same pockets, two predictor architectures, three mechanism '
             'variants', 'Independently computed Vina Dock score'],
            ['RQ2', 'Within-generator repeated-measures experiment (positive control)',
             'Best-of-N Vina-score reranking vs. unguided sampling, same pockets and independent evaluation pass',
             'Independently computed Vina Dock score, plus quality-degradation checks'],
            ['RQ3', 'Observational / diagnostic (no experimental manipulation of the generator)',
             'Gradient magnitude as a function of atomic distance from the pocket; factorial atom-removal ablation',
             'Spearman correlation; per-condition mean |Δ prediction|, target-clustered CI'],
            ['RQ4', 'Comparative benchmarking (representation ablation / factorial feature blocks)',
             'Typed-interaction features vs. ligand-only / pocket-only / interaction-plus-size feature blocks',
             'Target-clustered, BH-corrected R²'],
            ['RQ5', 'Quasi-experimental manipulation (within-target label permutation, a randomisation-test design)',
             'Real labels vs. within-target-permuted labels, same architecture, hyperparameters and split',
             'R² under permutation vs. R² under real labels'],
            ['RQ6', 'Observational statistical power analysis, derived from repeat-measurement label noise and '
             'measured between-seed model variance', 'Noise-imposed ceiling vs. ligand-only floor vs. average CI '
             'width', 'R² window width vs. average bootstrap CI width'],
        ],
        widths=[0.5, 2.0, 2.3, 1.7],
    )

    h2('1.6  Scope, Significance, and Contributions')
    para('The thesis this chapter introduces is not organised around a single positive claim, but around a '
         '*falsification design*: guidance is implemented and tested in good faith, and when it fails, the same '
         'statistical and diagnostic discipline is turned on the failure itself rather than treated as an endpoint. '
         'Null and negative results are reported here as legitimate scientific outcomes, not as a shortfall to be '
         'minimised, consistent with the course’s own guidance that transparent reporting includes negative, null, '
         'or contradictory findings together with their limitations.')
    para('Within that design, five claims are judged, after an eight-round literature search, to be defensible as '
         'novel at the time of writing: (1) a factorial ablation that separates the *quantity* of protein atoms '
         'present from the *identity* of the interface they form, which no attribution-based study examined was '
         'found to do; (2) a class-prior decomposition of apparent affinity-model performance via within-target '
         'y-randomisation, replicated for the deployed 3-D equivariant architecture itself and not only for a '
         'descriptor baseline; (3) a task-level statistical power analysis showing the resolvable performance '
         'window is narrower than one confidence interval on this benchmark; (4) the explicit connection from an '
         'information-content audit of the guidance predictor to the observed guidance failure, which situates the '
         'guidance literature and the affinity-prediction-reliability literature against each other directly; and '
         '(5) ternary (1.58-bit) quantization-aware training of a 3-D equivariant affinity network, validated not '
         'only on accuracy (indistinguishable from full precision at R² = 0.3437 ± 0.0377 vs. 0.3420) but on '
         'geometric equivariance under rigid motion, where the measured deviation sits at the floating-point ULP '
         'floor itself. One further claim, concerning the best-of-N reranking mechanism used for RQ2, was narrowed '
         'after a directed literature search found that best-of-N/best-of-K selection for diffusion SBDD had '
         'already been proposed as a generation technique in its own right @{bokdiff2025}; the defensible '
         'contribution of RQ2 is not the selection mechanism itself but its use as a *diagnostic instrument* that '
         'isolates signal quality from guidance mechanism by direct comparison on the same generator and pockets, '
         'a use not reported in that source.')
    para('Beyond these five, several findings replicate and sharpen results already present in the literature — '
         'most prominently that ligand-only descriptors are statistically indistinguishable from a 3-D '
         'structure-aware model on this leakage-controlled split, consistent with earlier reports that '
         'structure-based machine-learned scoring functions do not reliably outperform ligand-only or '
         'bias-only baselines @{volkov2022}. These are reported as replications, not as new contributions, in '
         'keeping with the course’s distinction between novelty and rigor as separate criteria for a defensible '
         'research problem.')

    h2('1.7  Organization of the Remaining Chapters')
    para('Chapter II reviews the literature underlying each research question above, organised by theme rather '
         'than by chapter of the underlying thesis, and closes with a conceptual framework that positions this '
         'work against the closest prior studies. The methodology (data, models, guidance mechanisms, statistical '
         'protocol and integrity checks), the full experimental results for each research question, and the '
         'cross-cutting discussion and conclusion are reported in the companion thesis chapters and are outside '
         'the scope of this Week 5 draft, which covers Chapters I and II only, as specified by the assignment.')

    # ================================================================================================ CHAPTER II
    h1('CHAPTER II — LITERATURE REVIEW')
    para('This chapter reviews the literature that research questions RQ1–RQ6 build on. It proceeds thematically: '
         'from the generative models that are guided (2.1), through the guidance mechanisms themselves (2.2), to '
         'the predictors that supply a guidance signal and the two literatures that determine whether such a '
         'predictor can be trusted — leakage and shortcut learning (2.4) and the reliability of docking-based '
         'evaluation (2.5) — then to uncertainty quantification (2.6), model compression and equivariance (2.7), '
         'standard QSAR controls this project adapts (2.8), and synthesizability estimation (2.9). Section 2.10 '
         'draws these threads into a conceptual framework and states the research gap directly. Statements are '
         'limited to what the cited sources report; where a source’s full methods were not read, this is stated '
         'explicitly rather than implied.')

    h2('2.1  Deep Generative Models for Structure-Based Drug Design')
    para('The generation of ligands conditioned on a protein pocket has been approached with several model '
         'families. Non-autoregressive models based on denoising diffusion @{ho2020}, which learn to reverse a '
         'gradual noising process, have displaced earlier voxel-grid and autoregressive atom-placement approaches '
         'for this task. TargetDiff @{guan2023} applies an equivariant diffusion model to the joint distribution '
         'of ligand atom coordinates and atom types conditioned on a protein pocket, and additionally uses the '
         'learned representation for affinity prediction; it is trained on CrossDocked2020 @{francoeur2020}, a set '
         'of 22.5 million docked poses of ligands cross-docked into similar binding pockets of the Protein Data '
         'Bank — poses that are themselves docking outputs rather than experimental complexes, a property whose '
         'consequences for evaluation are returned to in Section 2.5.')
    para('Equivariant graph networks @{satorras2021} provide the architectural basis for such models: node '
         'features are updated from invariant quantities such as inter-atomic distances, and coordinates are '
         'updated along relative position vectors, so that rotating or translating a complex rotates or translates '
         'the output accordingly. The same family of networks supplies the affinity predictor used for guidance in '
         'this project. A related strand of work injects interaction knowledge into 3-D generators directly: '
         'DeepICL conditions an autoregressive generator on atom-level interaction labels (hydrogen bond, salt '
         'bridge, hydrophobic, π-stacking) extracted with the protein–ligand interaction profiler PLIP '
         '@{zhung2024,salentin2015}; InterDiff conditions a diffusion model on interaction prompts via '
         'cross-attention and an auxiliary classification loss inside the denoiser @{wu2024}; BInD generates '
         'non-covalent interactions jointly with atoms and bonds under its own loss @{lee2025bind}; and DPDiff '
         'extracts interaction priors from pretrained affinity networks during sampling, reporting that a '
         'geometry-based prior pretrained on accurate poses *degraded* generation when queried on noisy '
         'intermediate diffusion states, because of a mismatch between training and sampling-time inputs '
         '@{huang2026} — an observation directly relevant to the gradient-guidance mismatch examined under RQ1 and '
         'RQ3.')

    h2('2.2  Guidance of Diffusion Models for SBDD')
    para('Classifier guidance modifies the reverse diffusion process with the gradient of an auxiliary predictor '
         'evaluated on noisy or partially denoised inputs, trading diversity for fidelity to a desired property '
         '@{dhariwal2021}. Two design decisions distinguish implementations of this idea for molecular generation. '
         'First, the guidance predictor may be trained directly on noisy inputs at every diffusion timestep, or '
         'trained on clean structures and evaluated only on the model’s clean-data estimate; the latter is cheaper '
         'but exposes the predictor to inputs that differ from its training distribution, particularly early in '
         'sampling. Second, guidance may act on continuous coordinates alone or also on discrete atom types, the '
         'latter requiring a differentiable relaxation. Both decisions matter directly for RQ1 and RQ3.')
    para('The closest prior work to RQ1 is BADGER, whose classifier-guidance variant applies the gradient of a '
         'separately trained affinity predictor to a frozen diffusion SBDD generator in essentially the same '
         'plug-and-play arrangement used here, reporting up to a 60% improvement in binding affinity '
         '@{jian2026badger}; a companion classifier-free variant instead folds affinity conditioning into the '
         'generator’s own training, which the frozen-generator design of this project does not attempt. DiffGui '
         'integrates bond diffusion with affinity and drug-likeness guidance trained jointly with the generator '
         'rather than through a separately frozen predictor, and reports improved binding affinity together with a '
         'small set of wet-laboratory assays @{hu2025diffgui}. Within the abstracts and summary materials examined '
         'for this review (Appendix B of the underlying thesis records the search strategy and its limits), '
         'neither source states whether its reported improvement was evaluated on a target-disjoint split, tested '
         'with correction for multiple comparisons, or checked against a positive/negative control or a '
         'quality-degradation criterion of the kind RQ1 and RQ2 apply; this does not claim those checks are absent '
         'from their full methods, only that this review did not confirm their presence.')
    para('A mechanistically different route to guidance avoids gradients altogether: best-of-N or best-of-K '
         'selection scores an oversampled, unguided pool of candidates with an external value function and keeps '
         'only the top-ranked subset. BoKDiff applies this directly to diffusion SBDD, ranking candidates generated '
         'by DecompDiff on a combined QED, SA and Vina score, without any fine-tuning, and explicitly describes '
         'itself as the first application of best-of-K alignment and best-of-N sampling to structure-based drug '
         'design @{bokdiff2025}. This is mechanistically identical to the selection procedure used for RQ2 in this '
         'project — oversample without guidance, rank by an external score, evaluate the kept subset — so the '
         'contribution of RQ2 cannot be the selection mechanism itself, which is prior art; it is the use of that '
         'mechanism as a controlled comparison against gradient guidance *on the same generator and pockets*, '
         'framed explicitly as a positive control for an evaluation harness rather than as a proposed generation '
         'technique, and decomposed against a three-mechanism account (uninformative signal, misaligned gradient, '
         'gameable objective) that @{bokdiff2025} does not address. An older, 2-D precedent for guidance by a '
         'genuine (non-learned-surrogate) scoring signal is reinforcement-learning-based generation guided directly '
         'by docking scores, reported to succeed for the REINVENT framework on a GPCR case study @{thomas2021}; '
         'this result, outside the diffusion-SBDD literature, is part of the basis for treating a positive RQ2 '
         'result as the a priori more likely outcome before the experiment was run.')

    h2('2.3  Learned Binding-Affinity Prediction')
    h3('2.3.1  Docking scores and their biases')
    para('AutoDock Vina combines an empirical scoring function with efficient conformational search and has been '
         'extended with additional features and Python bindings in later releases @{trott2010,eberhardt2021}. Its '
         'score, reported in kcal/mol, is used ubiquitously both to evaluate generated molecules and as a label or '
         'baseline for machine-learned scoring, but its limitations are documented: a systematic analysis of '
         'docking successes and failures on the DUD-E benchmark reports a bias toward compounds of higher '
         'molecular weight @{xu2022}, and re-parameterisations such as Vinardo were introduced because fitting '
         'empirical terms by linear regression to experimental affinities does not by itself yield a function with '
         'optimal docking capability @{quiroga2016}. This bias is directly relevant to interpreting the RQ2 result, '
         'where reranked molecules are both better-scoring and larger.')
    h3('2.3.2  Learned scoring functions')
    para('Machine-learned scoring functions built on graph and convolutional architectures include GIGN, which '
         'unifies covalent and non-covalent interactions in one rotation- and translation-invariant message-passing '
         'layer @{yang2023gign}; PIGNet, which predicts physics-informed pairwise interaction terms whose '
         'parameters are produced by a neural network @{moon2022}; Interformer, which combines a graph-transformer '
         'with an interaction-aware mixture-density network @{lai2024}; and PLANET, trained jointly on affinity, a '
         'protein–ligand contact map and a ligand distance matrix @{zhang2024planet}. The affinity predictor used '
         'for guidance in this project is an equivariant graph network in the same architectural family as GIGN.')

    h2('2.4  Data Leakage and Shortcut Learning in Affinity Prediction')
    para('A substantial part of the recent literature concerns whether reported affinity-model performance is '
         'genuine. Using modular graph networks, @{volkov2022} report that an explicit description of '
         'protein–ligand non-covalent interactions gives no advantage over ligand- or protein-only descriptors, '
         'and that simple nearest-neighbour baselines on the training set already perform well — evidence '
         'interpreted as memorisation dominating true learning, and directly motivating RQ4 and the ligand-only '
         'comparisons used throughout this project. Retraining top models on PDBbind CleanSplit, which removes '
         'structure-based train–test leakage and within-training-set redundancy, caused benchmark performance to '
         'drop substantially @{graber2025}. The Leak Proof PDBBind data set reorganises PDBbind so that proteins '
         'and ligands with high sequence or structural similarity never cross the split @{li2026lp}; the '
         'leakage-safe, target-level split used throughout this project (Section 1.4) follows the spirit of this '
         'procedure, adapted to CrossDocked2020, and is not the published data set itself. Complementary evidence '
         'comes from a study in which injecting additional geometric pocket-surface descriptors gave no '
         'statistically significant improvement on a leakage-controlled split @{balci2026}, and from a scoring '
         'function whose feature attributions were compared against a distance-based interaction profiler to test '
         'whether it had learned genuine interactions rather than dataset biases @{scantlebury2023}.')

    h2('2.5  Evaluating Generated Molecules')
    para('The reliability of docking scores as the success criterion for generative SBDD models has been '
         'questioned directly. @{gao2025} report that Vina scores can be inflated simply by increasing the number '
         'of atoms in a molecule, that several current SBDD models with high Vina scores fall short on measures of '
         'practical usability, and — in the same study — that some models use trained Vina predictors to guide '
         'sampling in a way that prioritises “hacking” the Vina score over generating reliable, effective '
         'molecules, and that selecting checkpoints on the test set risks data leakage. Physical plausibility is a '
         'separate axis of evaluation: PoseBusters, a suite of chemical and geometric checks (stereochemistry, '
         'bond lengths, protein–ligand clashes), finds that several deep-learning docking and generation methods '
         'frequently produce physically implausible poses even when root-mean-square deviation from a native pose '
         'is low @{buttenschoen2024}. A related study of interaction recovery reports that ignoring interaction '
         'fingerprints can overestimate model performance, particularly for methods that fail to recapitulate key '
         'interactions even while matching pose geometry @{errington2025}. Together, these works motivate the '
         'multi-criterion evaluation used for RQ1 and RQ2 — a docking-score improvement is never interpreted '
         'without a paired ligand-efficiency and physical-validity check.')

    h2('2.6  Uncertainty Quantification for Learned Affinity Predictors')
    para('A companion strand of this project’s broader work (outside RQ1–RQ6, but relevant to how guidance '
         'confidence could in principle be gated) evaluated six uncertainty-quantification methods — Monte Carlo '
         'dropout, deep ensembles, heteroscedastic Gaussian negative log-likelihood, deep evidential regression, a '
         'corrected evidential variant, and split-conformal prediction — against a pre-registered coverage test. A '
         'close prior-art check found that comparing uncertainty-quantification methods for protein–ligand affinity '
         'is not itself new: @{rayka2025} compare five UQ methods (deep ensemble, MC dropout, Laplace '
         'approximation, Bayes-by-backprop, evidential) on Leak-Proof PDBBind, overlapping substantially with this '
         'project’s own method list and dataset philosophy, and @{parks2020,rayka2024} report conformal prediction '
         'for binding affinity with good calibration, predating this project’s own positive conformal-prediction '
         'result. What remains open after that comparison is narrower: those studies use a feed-forward network on '
         'fixed ECIF-style descriptors, whereas this project evaluates the same question on a 3-D equivariant graph '
         'network under a target-clustered, multiply corrected, multi-seed-replicated inference framework rather '
         'than descriptive calibration metrics alone — and the strongest surviving claim from that comparison is a '
         'methodological one: an evidential-regression effect significant at one random seed (p = 0.042) '
         'disappeared across three seeds, with between-seed standard deviation in R² (0.038) exceeding the claimed '
         'effect size, a replication-failure pattern that applies to the broader single-seed UQ literature in this '
         'domain, including the sources above.')

    h2('2.7  Model Compression and the Equivariance Question')
    para('Reducing a trained network’s weights to a small discrete alphabet (quantization-aware training) is an '
         'established technique for deploying neural networks under tight memory or compute budgets; applying it '
         'to an equivariant geometric network raises a question not posed by quantizing a standard convolutional or '
         'transformer architecture: whether discretising weights can break the network’s equivariance to rotation '
         'and translation, a property the architecture is specifically designed to guarantee. @{zhou2026} examine '
         'this directly for an 8-bit SO(3)-equivariant attention transformer trained on QM9 and rMD17, show that '
         'naive low-bit quantization of equivariant *vector* features damages equivariance, introduce a Local '
         'Equivariance Error metric to track it, and report that preserving equivariance under quantization required '
         'three separate architectural interventions. This project’s own architecture updates coordinates through a '
         'learned *scalar* multiplying a relative position vector, rather than quantizing an equivariant vector '
         'feature directly, which is the specific structural difference this project’s own equivariance check was '
         'designed, before measurement, to exploit; the resulting prediction is that ternary quantization of this '
         'architecture should be equivariance-safe essentially by construction, in a way the architecture examined '
         'by @{zhou2026} is not.')

    h2('2.8  QSAR Controls: y-Randomisation, Activity Cliffs, and Matched Molecular Pairs')
    para('y-Randomisation — refitting a model after permuting the response variable and confirming performance '
         'collapses to chance — is a long-standing standard control against chance correlation in QSAR/QSPR '
         'modelling @{rucker2007}. Permuting labels *within* each target rather than globally adapts the same idea '
         'into a class-prior decomposition: it destroys within-target structure–activity relationships while '
         'preserving the typical affinity range associated with a target’s chemotype, isolating exactly the '
         'quantity RQ5 measures. Activity cliffs — pairs of structurally similar ligands with disproportionately '
         'different measured affinity — are a standard stress test for whether a model has learned local '
         'structure–activity relationships rather than coarse, class-level trends; the MoleculeACE benchmark '
         'reports that across 24 machine-learning approaches on 30 targets, descriptor-based models outperform more '
         'complex deep-learning models specifically on cliff compounds @{vantilborg2022}, a pattern this project’s '
         'own ligand-only-versus-structure comparisons echo. Matched molecular pairs, which isolate the effect of a '
         'single substituent change, are a related and stricter test of local structure–activity capture; '
         '@{kwapien2022} report that additive structure–activity relationships are the easiest case for any model, '
         'machine-learned or otherwise, to capture — a baseline against which this project’s own matched-pair sign '
         'agreement result (both a structure-aware ensemble and a ligand-only ridge model fall at or below chance '
         'agreement with the true direction of a substituent’s effect) is a notably severe negative finding.')

    h2('2.9  Estimating Synthetic Accessibility')
    para('Two synthesizability estimates are in common use. The synthetic accessibility (SA) score combines '
         'fragment contributions derived from roughly one million PubChem molecules with a complexity penalty, '
         'mapping a molecule to a value between 1 (easy) and 10 (very difficult) @{ertl2009}. The retrosynthetic '
         'accessibility score (RAscore) is a machine-learned classifier trained to predict whether the '
         'computer-aided synthesis-planning tool AiZynthFinder can find a route for a compound, reported to be at '
         'least 4500 times faster than the underlying planning search @{thakkar2021}. Both are used in this project '
         'purely as independent evaluation metrics, computed after generation and never as training targets for the '
         'guidance predictor itself in the routes examined by RQ1 and RQ2; neither is a guarantee that a synthesis '
         'route actually exists.')

    h2('2.10  Conceptual Framework and Positioning')
    para('The results summarised under RQ1–RQ6 are organised around a single explanatory framework with three, not '
         'mutually exclusive, candidate mechanisms for why gradient-based guidance might fail to produce an '
         'independently detectable improvement: **M1**, the guidance signal itself carries little usable '
         'information about the true binding interface (supported by RQ3’s gradient-field analysis and RQ4’s '
         'representation audit, and by RQ2’s contrast — the same generator responds when the *signal* is replaced '
         'with a genuine one, even though the *gradient mechanism* was never used in RQ2); **M2**, the signal is '
         'informative but the specific gradient-guidance mechanism misapplies it (tested directly by varying '
         'predictor architecture and guidance mechanism under RQ1, with no configuration rescuing the effect); and '
         '**M3**, the guidance objective can be satisfied in a way that does not correspond to a genuine '
         'improvement, for instance by exploiting molecular size (the confound explicitly measured and disclosed, '
         'not hidden, in RQ2’s own result). Table 2.1 positions the closest prior studies against this framework '
         'and against the specific evaluation safeguards (leakage-safe split, corrected significance testing, '
         'explicit control, quality-degradation check) that RQ1 and RQ2 apply.')
    table(
        ['Work', '3-D diffusion SBDD', 'Separate, learned guidance predictor', 'Explicit positive/negative control',
         'Leakage-safe, target-disjoint split', 'Multiple-comparison correction'],
        [
            ['DeepICL @{zhung2024}', 'No (autoregressive)', 'No (conditioning, not guidance)', 'Not the focus, as reviewed',
             'Not the focus, as reviewed', 'Not the focus, as reviewed'],
            ['InterDiff @{wu2024}', 'Yes', 'No (in-denoiser conditioning)', 'Not the focus, as reviewed',
             'Not the focus, as reviewed', 'Not the focus, as reviewed'],
            ['BADGER @{jian2026badger}', 'Yes', 'Yes (classifier-guidance variant)',
             'Not stated in materials reviewed', 'Not stated in materials reviewed', 'Not stated in materials reviewed'],
            ['DiffGui @{hu2025diffgui}', 'Yes', 'No (jointly trained with generator)',
             'Not stated; includes wet-lab validation', 'Not stated in materials reviewed', 'Not stated in materials reviewed'],
            ['BoKDiff @{bokdiff2025}', 'Yes', 'No (external scoring, no gradient)',
             'No gradient-guidance comparison reported', 'Not the focus, as reviewed', 'Not the focus, as reviewed'],
            ['This research (RQ1–RQ6)', 'Yes (frozen TargetDiff)', 'Yes (EGNN family; also a non-gradient Vina control)',
             'Yes — RQ2 is an explicit, pre-registered positive control', 'Yes (LP-split-style adaptation)',
             'Yes (Benjamini–Hochberg across all paired comparisons)'],
        ],
        widths=[1.1, 1.0, 1.4, 1.5, 1.1, 1.1],
    )
    para('What was not found, within the limits of the searches recorded in the underlying thesis’s methodology '
         'appendix, is a study that evaluates learned gradient-based affinity guidance for a 3-D diffusion SBDD '
         'generator together with an explicit positive control built from a genuine, non-surrogate signal on the '
         'same generator and pockets, under a leakage-safe, statistically corrected protocol, and that uses the '
         'contrast between the two to adjudicate among the M1/M2/M3 mechanisms above. That is the gap RQ1 and RQ2 '
         'jointly close; RQ3–RQ6 supply the mechanism-level evidence needed to decide which of the three '
         'explanations the overall null result actually supports, rather than leaving a bare null result '
         'unexplained.')

    h2('2.11  Summary and Link to Chapter III')
    para('This chapter has reviewed five bodies of work — generative SBDD models and their guidance, affinity '
         'prediction and its reliability threats, generated-molecule evaluation, uncertainty quantification, and '
         'model compression under equivariance constraints — and has shown that no source examined combines the '
         'specific safeguards RQ1 and RQ2 jointly apply. Chapter III, outside the scope of this Week 5 draft, '
         'describes in full the data, frozen generator, predictor architectures, guidance mechanisms, statistical '
         'protocol (target-clustered bootstrap, Benjamini–Hochberg correction, pre-registered decision bands) and '
         'integrity checks (leakage audits, split-artefact checks) that operationalise the research design '
         'summarised in Table 1.2.')


# ================================================================================================== BUILD + SAVE
DRY = True
build_body()

order = ordered_keys(sorted(set(CITED)))
NUMBERING = {k: i + 1 for i, k in enumerate(order)}

DRY = False
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.left_margin, sec.right_margin = Inches(1.25), Inches(1.0)
sec.top_margin, sec.bottom_margin = Inches(1.0), Inches(1.0)
st = doc.styles['Normal']
st.font.name = FONT; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

build_body()

# ---------------------------------------------------------------------------------------------- references
h1('REFERENCES')
para('Sources are numbered sequentially in alphabetical order by first author’s surname and cited in text by '
     'number, following the project’s existing IEEE-style convention. Entries marked with an asterisk (*) are '
     'cited from this project’s own literature-search log by author, year and venue, and have not been '
     'individually re-verified against the original source for this draft; all other entries were checked '
     'against the primary source or its PubMed record during the underlying project (see refs.py in the '
     'project repository for the full verification ledger).', first=None)

for k in order:
    n = NUMBERING[k]
    star = ' *' if REFS[k]['v'] == 'kn' else ''
    text = f'[{n}] ' + REFS[k]['t'] + star
    p = doc.add_paragraph(); add_runs(p, text, 11)
    fmt(p, WD_ALIGN_PARAGRAPH.LEFT, 1.2, first=None, after=6)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)

doc.save(OUT_DOCX)
print('wrote', OUT_DOCX, '--', len(order), 'references cited')


def to_pdf(path):
    outdir = os.path.dirname(path)
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', outdir, path],
                    check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600)
    return path[:-5] + '.pdf'


pdf = to_pdf(OUT_DOCX)
print('wrote', pdf)
